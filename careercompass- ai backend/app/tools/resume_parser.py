"""Secure resume upload handling + text extraction (PDF / DOCX / TXT)."""
import io
import zipfile

from fastapi import HTTPException

from app.core.config import settings

ALLOWED = {".pdf", ".docx", ".txt"}


def _ext(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def validate_upload(filename: str, data: bytes) -> str:
    """Validate extension, size and magic bytes. Returns the extension."""
    ext = _ext(filename or "")
    if ext not in ALLOWED:
        raise HTTPException(415, f"Unsupported file type. Allowed: {', '.join(sorted(ALLOWED))}")
    if not data:
        raise HTTPException(400, "Empty file")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, f"File too large (max {settings.max_upload_mb} MB)")
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise HTTPException(415, "File content is not a valid PDF")
    if ext == ".docx" and not data.startswith(b"PK\x03\x04"):
        raise HTTPException(415, "File content is not a valid DOCX")
    return ext


def extract_text(filename: str, data: bytes) -> str:
    ext = validate_upload(filename, data)
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(422, "Encrypted PDFs are not supported")
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif ext == ".docx":
            from docx import Document

            with zipfile.ZipFile(io.BytesIO(data)) as z:  # zip-bomb guard
                if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                    raise HTTPException(413, "Document expands to an unreasonable size")
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="ignore")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Could not read this file. Is it corrupted?")
    text = text.replace("\x00", " ").strip()
    if len(text) < 50:
        raise HTTPException(422, "No readable text found (scanned image PDFs are not supported)")
    return text
